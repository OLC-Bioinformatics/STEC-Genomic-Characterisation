#!/usr/bin/env python3
"""
Sipprverse subtype-level StxDB validation comparison
====================================================

Compares parsed Sipprverse subtype hits against the collapsed subtype-level
StxDB ground truth and generates an Excel workbook with Summary,
Subtype_Details, Ground_Truth, and Metrics sheets.

Special functional-status rule
------------------------------
Two stx-like targets are retained in the curated validation table but are not
functional toxins. They are excluded from the expected-positive set:

* 2025-SEQ-1796: Stx2a is interrupted by a transposon.
* SRR18191635: Stx2b contains a frameshift.

If Sipprverse reports these subtype/genome combinations, they are classified
as false positives. If they are not reported, they do not create false
negatives. The functional Stx1a and Stx2d targets in those genomes remain
normal expected positives.

Example
-------
python sipprverse_subtype_comparison_pipeline_functional_rule.py \\
  --groundtruth StxDB_groundtruth_subtype_level.xlsx \\
  --sipprverse sipprverse_0.98_parsed_file.csv \\
  --output sipprverse_subtype_comparison_0.98.xlsx \\
  --cutoff 98
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path
from typing import Dict, List, Optional, Set, Tuple

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

TOTAL_STX_SUBTYPES = 19
KNOWN_SUBTYPE_ORDER = [
    "Stx1a", "Stx1c", "Stx1d", "Stx1e",
    "Stx2a", "Stx2b", "Stx2c", "Stx2d", "Stx2e", "Stx2f", "Stx2g",
    "Stx2h", "Stx2i", "Stx2j", "Stx2k", "Stx2l", "Stx2m", "Stx2n", "Stx2o",
]

# Genome-specific stx-like targets that must not be accepted as functional TPs.
NONFUNCTIONAL_SUBTYPES_BY_GENOME: Dict[str, Set[str]] = {
    "2025-SEQ-1796": {"Stx2a"},
    "SRR18191635": {"Stx2b"},
}

SUBTYPE_RE = re.compile(r"Stx[12][a-z]", re.IGNORECASE)


def normalize_subtype(value: object) -> Optional[str]:
    """Extract a subtype and normalize it to the form Stx1a or Stx2d."""
    if value is None or pd.isna(value):
        return None
    match = SUBTYPE_RE.search(str(value))
    if not match:
        return None
    raw = match.group(0)
    return "Stx" + raw[3].upper() + raw[4].lower()


def split_values(value: object) -> List[str]:
    """Split semicolon-, comma-, or newline-separated spreadsheet values."""
    if value is None or pd.isna(value):
        return []
    return [x.strip() for x in re.split(r"[;,\n]+", str(value)) if x.strip()]


def subtype_sort_key(subtype: str) -> Tuple[int, int, str]:
    if subtype in KNOWN_SUBTYPE_ORDER:
        return (0, KNOWN_SUBTYPE_ORDER.index(subtype), subtype)
    return (1, 999, subtype)


def join_subtypes(values: Set[str]) -> str:
    return "; ".join(sorted(values, key=subtype_sort_key))


def safe_divide(numerator: int, denominator: int) -> float:
    return numerator / denominator if denominator else 0.0


def load_groundtruth(path: Path) -> pd.DataFrame:
    """Read and validate the collapsed subtype-level ground-truth workbook."""
    df = pd.read_excel(path)
    required = {"Fastq", "Expected_Subtypes"}
    missing = required - set(df.columns)
    if missing:
        raise ValueError(
            f"Ground-truth file is missing required column(s): {sorted(missing)}"
        )

    df = df.copy()
    df["Genome"] = df["Fastq"].astype(str).str.strip()
    if df["Genome"].duplicated().any():
        duplicates = sorted(df.loc[df["Genome"].duplicated(False), "Genome"].unique())
        raise ValueError(f"Duplicate genomes in ground truth: {duplicates[:10]}")
    return df


def load_sipprverse(path: Path, cutoff: float) -> pd.DataFrame:
    """Read parsed Sipprverse hits and retain hits meeting the identity cutoff."""
    df = pd.read_csv(path)
    required = {"Genome", "Sipprverse_Hit", "Subtype", "Percent_Identity"}
    missing = required - set(df.columns)
    if missing:
        raise ValueError(
            f"Sipprverse parsed file is missing required column(s): {sorted(missing)}"
        )

    df = df.copy()
    df["Genome"] = df["Genome"].astype(str).str.strip()
    df["Percent_Identity"] = pd.to_numeric(df["Percent_Identity"], errors="coerce")
    df["Normalized_Subtype"] = df["Subtype"].apply(normalize_subtype)
    df = df[
        df["Normalized_Subtype"].notna()
        & df["Percent_Identity"].notna()
        & (df["Percent_Identity"] >= cutoff)
    ].copy()
    return df


def compare(
    groundtruth: pd.DataFrame,
    hits: pd.DataFrame,
) -> Tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Perform subtype-level classification for every ground-truth genome."""
    hits_by_genome = {g: group.copy() for g, group in hits.groupby("Genome")}

    summary_rows = []
    detail_rows = []

    for _, gt_row in groundtruth.iterrows():
        genome = gt_row["Genome"]

        curated_expected = {
            subtype
            for item in split_values(gt_row["Expected_Subtypes"])
            if (subtype := normalize_subtype(item)) is not None
        }

        nonfunctional = NONFUNCTIONAL_SUBTYPES_BY_GENOME.get(genome, set())
        # Only remove a special subtype from the positive set. It remains
        # documented in the all-curated and non-functional columns.
        functional_expected = curated_expected - nonfunctional

        genome_hits = hits_by_genome.get(genome, pd.DataFrame(columns=hits.columns))
        detected = set(genome_hits.get("Normalized_Subtype", pd.Series(dtype=str)).dropna())

        tp = detected & functional_expected
        fp = detected - functional_expected
        fn = functional_expected - detected
        tn = TOTAL_STX_SUBTYPES - len(tp) - len(fp) - len(fn)
        if tn < 0:
            raise ValueError(
                f"Negative TN for {genome}. Check subtype universe and parsed calls."
            )

        detected_nonfunctional = detected & nonfunctional

        summary_rows.append({
            "Genome": genome,
            "Expected_Operons": gt_row.get("Expected_Operons", ""),
            "Expected_All_Curated_Subtypes": join_subtypes(curated_expected),
            "Expected_Functional_Subtypes": join_subtypes(functional_expected),
            "Nonfunctional_Subtypes": join_subtypes(nonfunctional),
            "Detected_Subtypes": join_subtypes(detected),
            "Detected_Nonfunctional_Subtypes": join_subtypes(detected_nonfunctional),
            "TP": len(tp),
            "FP": len(fp),
            "FN": len(fn),
            "TN": tn,
            "Sensitivity": safe_divide(len(tp), len(tp) + len(fn)),
            "Specificity": safe_divide(tn, tn + len(fp)),
            "Precision": safe_divide(len(tp), len(tp) + len(fp)),
            "NPV": safe_divide(tn, tn + len(fn)),
            "Accuracy": safe_divide(len(tp) + tn, len(tp) + len(fp) + len(fn) + tn),
            "F1_Score": safe_divide(2 * len(tp), 2 * len(tp) + len(fp) + len(fn)),
        })

        for subtype in sorted(tp | fp | fn, key=subtype_sort_key):
            expected_flag = subtype in functional_expected
            detected_flag = subtype in detected
            classification = "TP" if subtype in tp else "FP" if subtype in fp else "FN"

            subtype_hits = genome_hits[
                genome_hits.get("Normalized_Subtype", pd.Series(dtype=str)) == subtype
            ].copy()
            if not subtype_hits.empty:
                best = subtype_hits.sort_values(
                    ["Percent_Identity", "Sipprverse_Hit"],
                    ascending=[False, True],
                ).iloc[0]
                best_hit = best["Sipprverse_Hit"]
                best_identity = best["Percent_Identity"]
            else:
                best_hit = ""
                best_identity = ""

            detail_rows.append({
                "Genome": genome,
                "Compared_Subtype": subtype,
                "Expected_Functional": expected_flag,
                "Detected": detected_flag,
                "Nonfunctional_Special_Case": subtype in nonfunctional,
                "Best_Sipprverse_Hit": best_hit,
                "Percent_Identity": best_identity,
                "Classification": classification,
            })

    summary = pd.DataFrame(summary_rows)
    details = pd.DataFrame(detail_rows)

    totals = {
        metric: int(summary[metric].sum()) for metric in ["TP", "FP", "FN", "TN"]
    }
    metrics = pd.DataFrame([
        {"Performance_Metric": "True Positives (TP)", "Value": totals["TP"]},
        {"Performance_Metric": "False Positives (FP)", "Value": totals["FP"]},
        {"Performance_Metric": "False Negatives (FN)", "Value": totals["FN"]},
        {"Performance_Metric": "True Negatives (TN)", "Value": totals["TN"]},
        {"Performance_Metric": "Sensitivity (Recall)", "Value": safe_divide(totals["TP"], totals["TP"] + totals["FN"])},
        {"Performance_Metric": "Specificity", "Value": safe_divide(totals["TN"], totals["TN"] + totals["FP"])},
        {"Performance_Metric": "Precision (PPV)", "Value": safe_divide(totals["TP"], totals["TP"] + totals["FP"])},
        {"Performance_Metric": "Negative Predictive Value (NPV)", "Value": safe_divide(totals["TN"], totals["TN"] + totals["FN"])},
        {"Performance_Metric": "Accuracy", "Value": safe_divide(totals["TP"] + totals["TN"], sum(totals.values()))},
        {"Performance_Metric": "F1 Score", "Value": safe_divide(2 * totals["TP"], 2 * totals["TP"] + totals["FP"] + totals["FN"])},
    ])

    return summary, details, metrics


def format_workbook(path: Path) -> None:
    """Apply simple thesis-friendly formatting: bold headers and no shading."""
    wb = load_workbook(path)
    thin = Side(style="thin", color="D9D9D9")

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin)

        for column_cells in ws.columns:
            letter = get_column_letter(column_cells[0].column)
            max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
            ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 42)

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    # Format metric values and per-genome metric columns as percentages.
    if "Metrics" in wb.sheetnames:
        ws = wb["Metrics"]
        for row in range(6, ws.max_row + 1):
            ws.cell(row=row, column=2).number_format = "0.00%"

    if "Summary" in wb.sheetnames:
        ws = wb["Summary"]
        header_to_col = {cell.value: cell.column for cell in ws[1]}
        for name in ["Sensitivity", "Specificity", "Precision", "NPV", "Accuracy", "F1_Score"]:
            col = header_to_col.get(name)
            if col:
                for row in range(2, ws.max_row + 1):
                    ws.cell(row=row, column=col).number_format = "0.00%"

    wb.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Compare parsed Sipprverse subtype hits with StxDB ground truth."
    )
    parser.add_argument("--groundtruth", required=True, type=Path)
    parser.add_argument("--sipprverse", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument(
        "--cutoff",
        type=float,
        default=98.0,
        help="Minimum percent identity to retain (default: 98.0).",
    )
    args = parser.parse_args()

    groundtruth = load_groundtruth(args.groundtruth)
    hits = load_sipprverse(args.sipprverse, args.cutoff)
    summary, details, metrics = compare(groundtruth, hits)

    # Report parsed genomes that do not occur in the benchmark table.
    benchmark_genomes = set(groundtruth["Genome"])
    unmatched = sorted(set(hits["Genome"]) - benchmark_genomes)
    unmatched_df = hits[hits["Genome"].isin(unmatched)].copy()

    args.output.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(args.output, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Summary", index=False)
        details.to_excel(writer, sheet_name="Subtype_Details", index=False)
        groundtruth.drop(columns=["Genome"], errors="ignore").to_excel(
            writer, sheet_name="Ground_Truth", index=False
        )
        metrics.to_excel(writer, sheet_name="Metrics", index=False)
        if not unmatched_df.empty:
            unmatched_df.to_excel(writer, sheet_name="Unmatched_Sipprverse", index=False)

    format_workbook(args.output)

    print(f"Saved: {args.output}")
    print(f"Ground-truth genomes: {len(groundtruth)}")
    print(f"Retained Sipprverse hit rows: {len(hits)}")
    print(f"Unmatched Sipprverse genomes: {len(unmatched)}")
    print(metrics.to_string(index=False))


if __name__ == "__main__":
    main()
