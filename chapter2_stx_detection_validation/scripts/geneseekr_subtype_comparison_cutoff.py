#!/usr/bin/env python3
"""
Compare parsed GeneSeekr subtype calls with the subtype-level ground truth,
applying a minimum percent_match cutoff to the subtype call.

Rationale (C. Carrillo): a BLAST hit covering only part of the reference operon
cannot support a subtype assignment. Such a hit shows the genome carries stx,
not which subtype it carries. Hits below the cutoff are therefore not counted as
subtype detections; they are reported separately.

The cutoff is applied to the GeneSeekr percent_match column, which is the
proportion of the reference operon covered by the alignment
(alignment_length / subject_length).

Same conventions as the uncut version:
- 19 recognised subtypes define the negative space;
- a non-functional subtype cannot be a TP, is an FP when detected, and is not
  an FN when absent.

Usage:
python geneseekr_subtype_comparison_cutoff.py \
  --groundtruth StxDB_groundtruth_subtype_level.xlsx \
  --geneseekr parsed_geneseekr_outputs_best_hit_per_region.csv \
  --output GeneSeekr_Subtype_Comparison_90cov.xlsx \
  --min-percent-match 90
"""

import argparse
import csv
import re

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

TOTAL_SUBTYPES = 19

SUBTYPE_ORDER = [
    "Stx1a", "Stx1c", "Stx1d", "Stx1e",
    "Stx2a", "Stx2b", "Stx2c", "Stx2d", "Stx2e",
    "Stx2f", "Stx2g", "Stx2h", "Stx2i", "Stx2j",
    "Stx2k", "Stx2l", "Stx2m", "Stx2n", "Stx2o",
]

NONFUNCTIONAL_SUBTYPES = {
    "2025-SEQ-1796": {"Stx2a"},
    "SRR18191635": {"Stx2b"},
}

SUBTYPE_PATTERN = re.compile(r"Stx[12][a-z]", re.IGNORECASE)


def normalize_subtype(value):
    if value is None:
        return None
    match = SUBTYPE_PATTERN.search(str(value))
    if not match:
        return None
    return "Stx%s%s" % (match.group(0)[3], match.group(0)[4].lower())


def split_values(value):
    if value is None:
        return []
    return [item.strip() for item in re.split(r"[;,\n]+", str(value)) if item.strip()]


def sort_subtypes(values):
    return sorted(values, key=lambda s: (SUBTYPE_ORDER.index(s) if s in SUBTYPE_ORDER else 999, s))


def join_subtypes(values):
    return "; ".join(sort_subtypes(values))


def safe_ratio(n, d):
    return n / d if d else 0.0


def numeric(value, default=0.0):
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def load_ground_truth(path):
    df = pd.read_excel(path, sheet_name="Operon_vs_Subtype")
    df.columns = [str(c).strip() for c in df.columns]
    required = {"Fastq", "Fasta", "Expected_Operons", "Expected_Subtypes"}
    missing = required - set(df.columns)
    if missing:
        raise ValueError("Operon_vs_Subtype is missing: " + ", ".join(sorted(missing)))

    records = []
    for _, row in df.iterrows():
        if pd.isna(row["Fasta"]):
            continue
        expected = set()
        for item in split_values(row.get("Expected_Subtypes")):
            s = normalize_subtype(item)
            if s:
                expected.add(s)
        records.append({
            "Fastq": str(row["Fastq"]).strip(),
            "Fasta": str(row["Fasta"]).strip(),
            "Expected_Operons": "" if pd.isna(row.get("Expected_Operons")) else row["Expected_Operons"],
            "Expected_Subtypes": expected,
            "Operon_Count": row.get("Operon_Count"),
        })
    return records


def load_geneseekr_hits(path):
    """Every best-hit-per-region row, with coverage of the reference operon."""
    hits = {}
    with open(path, newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        required = {"sample_name", "subject_id", "percent_match", "bit_score",
                    "evalue", "alignment_length", "subject_length"}
        missing = required - set(reader.fieldnames or [])
        if missing:
            raise ValueError("Parsed CSV missing: " + ", ".join(sorted(missing)))

        for row in reader:
            fasta = str(row["sample_name"]).strip()
            subtype = normalize_subtype(row.get("subject_id"))
            if not fasta or subtype is None:
                continue
            row["_subtype"] = subtype
            row["_percent_match"] = numeric(row.get("percent_match"))
            hits.setdefault(fasta, []).append(row)
    return hits


def best_hit(rows):
    return sorted(rows, key=lambda r: (-numeric(r.get("bit_score")),
                                       numeric(r.get("evalue"), float("inf")),
                                       -numeric(r.get("percent_match")),
                                       -numeric(r.get("alignment_length"))))[0]


THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)


def write_sheet(ws, headers, rows, widths):
    ws.append(headers)
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=len(headers)):
        for cell in row:
            cell.font = Font(name="Calibri", size=11, bold=(cell.row == 1))
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.border = BORDER
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--groundtruth", required=True)
    ap.add_argument("--geneseekr", required=True)
    ap.add_argument("--output", required=True)
    ap.add_argument("--min-percent-match", type=float, default=90.0)
    args = ap.parse_args()

    cutoff = args.min_percent_match
    ground_truth = load_ground_truth(args.groundtruth)
    hits_by_fasta = load_geneseekr_hits(args.geneseekr)

    known = {r["Fasta"] for r in ground_truth}
    unmatched = sorted(set(hits_by_fasta) - known)
    if unmatched:
        raise ValueError("GeneSeekr samples absent from ground truth: " + ", ".join(unmatched[:10]))

    summary_rows, detail_rows, gt_rows, excluded_rows = [], [], [], []
    tp = fp = fn = tn = 0

    for rec in ground_truth:
        fastq, fasta = rec["Fastq"], rec["Fasta"]
        nonfunctional = NONFUNCTIONAL_SUBTYPES.get(fastq, set())
        expected = set(rec["Expected_Subtypes"]) - nonfunctional

        genome_hits = hits_by_fasta.get(fasta, [])
        passing = [h for h in genome_hits if h["_percent_match"] >= cutoff]
        failing = [h for h in genome_hits if h["_percent_match"] < cutoff]

        detected = {h["_subtype"] for h in passing}
        dropped = {h["_subtype"] for h in failing} - detected

        for h in failing:
            excluded_rows.append([
                fastq, fasta, h["_subtype"], h.get("subject_id", ""),
                round(h["_percent_match"], 2),
                int(numeric(h.get("alignment_length"))), int(numeric(h.get("subject_length"))),
                "Yes" if h["_subtype"] in detected else "No",
                "Yes" if h["_subtype"] in expected else "No",
            ])

        true_pos = detected & expected
        false_pos = detected - expected
        false_neg = expected - detected
        true_neg = TOTAL_SUBTYPES - len(true_pos) - len(false_pos) - len(false_neg)
        if true_neg < 0:
            raise ValueError("Negative TN for %s" % fastq)

        tp += len(true_pos); fp += len(false_pos); fn += len(false_neg); tn += true_neg

        summary_rows.append([
            fastq, fasta, join_subtypes(expected), join_subtypes(detected),
            join_subtypes(dropped), len(true_pos), len(false_pos), len(false_neg), true_neg,
        ])
        gt_rows.append([
            fastq, fasta, rec["Expected_Operons"], join_subtypes(expected),
            rec["Operon_Count"], len(expected),
        ])

        for subtype in sort_subtypes(true_pos | false_pos | false_neg):
            sub_hits = [h for h in passing if h["_subtype"] == subtype]
            if sub_hits:
                b = best_hit(sub_hits)
                hit_id, pm = b.get("subject_id", ""), round(b["_percent_match"], 2)
            else:
                below = [h for h in failing if h["_subtype"] == subtype]
                if below:
                    b = best_hit(below)
                    hit_id = "%s (below cutoff)" % b.get("subject_id", "")
                    pm = round(b["_percent_match"], 2)
                else:
                    hit_id, pm = "", ""
            cls = "TP" if subtype in true_pos else ("FP" if subtype in false_pos else "FN")
            detail_rows.append([
                fastq, fasta, subtype,
                "True" if subtype in expected else "False",
                "True" if subtype in detected else "False",
                hit_id, pm, cls,
            ])

    metrics = {
        "TP": tp, "FP": fp, "FN": fn, "TN": tn,
        "Sensitivity": safe_ratio(tp, tp + fn),
        "Specificity": safe_ratio(tn, tn + fp),
        "Precision": safe_ratio(tp, tp + fp),
        "Accuracy": safe_ratio(tp + tn, tp + fp + fn + tn),
        "F1 Score": safe_ratio(2 * tp, 2 * tp + fp + fn),
    }

    wb = Workbook()
    ws = wb.active; ws.title = "Summary"
    write_sheet(ws,
                ["Fastq", "Fasta", "Expected_Subtypes", "Detected_Subtypes",
                 "Subtypes_below_cutoff", "TP", "FP", "FN", "TN"],
                summary_rows, [18, 18, 27, 27, 27, 9, 9, 9, 9])

    write_sheet(wb.create_sheet("Subtype_Details"),
                ["Fastq", "Fasta", "Compared_Subtype", "Expected", "Detected",
                 "Best_GeneSeekr_Hit", "Percent_Match", "Classification"],
                detail_rows, [18, 18, 18, 12, 12, 30, 16, 15])

    write_sheet(wb.create_sheet("Excluded_Hits"),
                ["Fastq", "Fasta", "Subtype", "GeneSeekr_Hit", "Percent_Match",
                 "Alignment_length", "Reference_length",
                 "Subtype_also_called_above_cutoff", "Subtype_expected"],
                excluded_rows, [18, 18, 10, 30, 16, 16, 16, 26, 16])

    write_sheet(wb.create_sheet("Ground_Truth"),
                ["Fastq", "Fasta", "Expected_Operons", "Expected_Subtypes",
                 "Operon_Count", "Unique_Subtype_Count"],
                gt_rows, [18, 18, 42, 26, 14, 20])

    mws = wb.create_sheet("Metrics")
    write_sheet(mws, ["Metric", "Value"],
                [["Minimum percent_match (%)", cutoff],
                 ["TP", metrics["TP"]], ["FP", metrics["FP"]],
                 ["FN", metrics["FN"]], ["TN", metrics["TN"]],
                 ["Sensitivity", metrics["Sensitivity"]],
                 ["Specificity", metrics["Specificity"]],
                 ["Precision", metrics["Precision"]],
                 ["Accuracy", metrics["Accuracy"]],
                 ["F1 Score", metrics["F1 Score"]]],
                [32, 16])
    for r in range(7, 12):
        mws.cell(r, 2).number_format = "0.00%"

    wb.save(args.output)

    print("cutoff: percent_match >= %.1f%%" % cutoff)
    for k in ["TP", "FP", "FN", "TN"]:
        print("  %-12s %d" % (k, metrics[k]))
    for k in ["Sensitivity", "Specificity", "Precision", "Accuracy", "F1 Score"]:
        print("  %-12s %.4f" % (k, metrics[k]))
    print("  excluded hits: %d" % len(excluded_rows))


if __name__ == "__main__":
    main()
